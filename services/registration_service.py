"""注册表服务（蓝图「飞升」§快速启动）：《注册CommWeb.xlsx》模板生成 / 解析 / 校验 /
受限 registration migration JSON 生成 / 审核后导入 site_views。

铁律：生成物是纯数据声明（与 site_views 同构），不是 SQL、不含可执行代码；
apply 只做 upsert（不删不覆盖清单外行），导入前重新校验（不信任客户端）。
组件口径与 CommWEB src/protocol/registry.ts 对账维护（两端测试同规则）。
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from store.site_repo import SiteRepo

SHEET_NAME = "视图注册"

COLUMNS = ["id", "type", "title", "icon", "navKind", "navGroup", "navOrder",
           "default", "sort", "whenCapability", "propsJson"]

ID_PATTERN = re.compile(r"^[a-z][a-z0-9-]*$")

# 组件口径（与 CommWEB registry.ts 同步维护）
VIEW_TYPES = [
    ("workspace.tabs", "工作区（多标签会话）"),
    ("tools.grid", "工具中心"),
    ("flows.list", "流程中心"),
    ("tasks.table", "任务列表"),
    ("data.browser", "数据浏览"),
    ("pipeline.studio", "流程工作台"),
    ("settings.keys", "密钥管理"),
    ("templates.manager", "识别模板管理"),
    ("ocr.studio", "识别工作台"),
    ("translate.studio", "翻译工作台"),
]
VIEW_TYPE_IDS = [t for t, _ in VIEW_TYPES]

ICON_WHITELIST = [
    "appstore-outlined", "unordered-list-outlined", "desktop-outlined", "node-index-outlined",
    "table-outlined", "profile-outlined", "partition-outlined", "cluster-outlined",
    "database-outlined", "file-text-outlined", "api-outlined", "translation-outlined",
]

NAV_KINDS = ["tab", "menu", "child", "header", "hidden"]

CAPABILITY_KEYS = ["has_pipelines", "has_files", "has_runs", "has_sse", "has_statuses"]

EXAMPLE_ROWS = [
    ["work", "workspace.tabs", "工作区", "desktop-outlined", "menu", None, None, "是", 10, None, None],
    ["work-tasks", "tasks.table", "任务记录", "unordered-list-outlined", "child", "work", None, None, None, None, None],
    ["keys", "settings.keys", "密钥管理", "api-outlined", "header", None, 5, None, None, "has_pipelines", None],
]


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _checksum(payload: Any) -> str:
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:16]


# ── 模板生成 ───────────────────────────────────────────────

def build_template() -> bytes:
    """生成《注册CommWeb.xlsx》：表头 + 示例行 + 枚举下拉约束 + 填写说明 sheet。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(COLUMNS)
    for row in EXAMPLE_ROWS:
        ws.append(row)

    def dv(formula: str, cols: list[str]) -> None:
        rule = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
        ws.add_data_validation(rule)
        for col in cols:
            rule.add(f"{col}2:{col}500")

    dv(f'"{",".join(VIEW_TYPE_IDS)}"', ["B"])          # type
    dv(f'"{",".join(ICON_WHITELIST)}"', ["D"])          # icon
    dv(f'"{",".join(NAV_KINDS)}"', ["E"])               # navKind
    dv(f'"{",".join(CAPABILITY_KEYS)}"', ["J"])         # whenCapability
    dv('"是,否"', ["H"])                                 # default

    guide = wb.create_sheet("填写说明")
    guide_rows = [
        ["列", "必填", "说明"],
        ["id", "是", "视图唯一标识，小写字母开头，仅小写字母/数字/连字符"],
        ["type", "是", "视图类型（下拉选择），见「组件清单」sheet"],
        ["title", "是", "导航/页签标题"],
        ["icon", "否", "图标白名单（下拉选择），留空用默认"],
        ["navKind", "否", "导航形态：tab（默认）/menu/child/header/hidden"],
        ["navGroup", "条件", "navKind=child 时必填，指向 menu 父项 id"],
        ["navOrder", "否", "header 项排序，小者在前"],
        ["default", "否", "落地页（填「是」，全表至多一个）"],
        ["sort", "否", "排序权重"],
        ["whenCapability", "否", "能力键：探测为真才显示"],
        ["propsJson", "否", '附加 props（JSON 对象），如 {"defaultLabel":"识别中心"}'],
    ]
    for row in guide_rows:
        guide.append(row)

    catalog = wb.create_sheet("组件清单")
    catalog.append(["视图类型", "名称", "支持槽位"])
    supports_slots = {"tools.grid", "flows.list", "tasks.table", "data.browser"}
    for type_id, name in VIEW_TYPES:
        catalog.append([type_id, name, "是" if type_id in supports_slots else "否"])

    for sheet in (ws, guide, catalog):
        for column_cells in sheet.columns:
            width = max(len(str(c.value or "")) for c in column_cells) + 4
            sheet.column_dimensions[column_cells[0].column_letter].width = min(width, 60)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ── 解析与校验 ─────────────────────────────────────────────

def _as_bool(value: Any) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip()
    if text in ("是", "真", "TRUE", "True", "1", "yes", "Y"):
        return True
    if text in ("否", "假", "FALSE", "False", "0", "no", "N"):
        return False
    return None


def parse_workbook(data: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"缺少工作表「{SHEET_NAME}」")
    ws = wb[SHEET_NAME]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header or [str(h).strip() for h in header[:len(COLUMNS)]] != COLUMNS:
        raise ValueError(f"表头不符：应为 {COLUMNS}")
    out: list[dict[str, Any]] = []
    for values in rows:
        if all(v is None or str(v).strip() == "" for v in values):
            continue  # 空行跳过
        cell = dict(zip(COLUMNS, values))
        out.append({
            "id": str(cell["id"]).strip() if cell["id"] is not None else "",
            "type": str(cell["type"]).strip() if cell["type"] is not None else "",
            "title": str(cell["title"]).strip() if cell["title"] is not None else "",
            "icon": str(cell["icon"]).strip() if cell["icon"] not in (None, "") else None,
            "navKind": str(cell["navKind"]).strip() if cell["navKind"] not in (None, "") else None,
            "navGroup": str(cell["navGroup"]).strip() if cell["navGroup"] not in (None, "") else None,
            "navOrder": int(cell["navOrder"]) if isinstance(cell["navOrder"], (int, float)) else None,
            "default": _as_bool(cell["default"]),
            "sort": int(cell["sort"]) if isinstance(cell["sort"], (int, float)) else None,
            "whenCapability": str(cell["whenCapability"]).strip() if cell["whenCapability"] not in (None, "") else None,
            "propsJson": str(cell["propsJson"]).strip() if cell["propsJson"] not in (None, "") else None,
        })
    return out


def validate_rows(rows: list[dict[str, Any]]) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    seen: dict[str, int] = {}
    menu_ids: set[str] = set()

    for i, row in enumerate(rows):
        at = f"第 {i + 2} 行（{row['id'] or '?'}）"
        rid = row["id"]
        if not rid:
            errors.append(f"{at}: id 必填")
        elif not ID_PATTERN.match(rid):
            errors.append(f"{at}: id 须为小写字母开头、仅含小写字母/数字/连字符")
        elif rid in seen:
            errors.append(f"{at}: id 重复（首见于第 {seen[rid] + 2} 行）")
        else:
            seen[rid] = i
        if row["type"] not in VIEW_TYPE_IDS:
            errors.append(f"{at}: 未知视图类型 {row['type']}")
        if not row["title"]:
            errors.append(f"{at}: title 必填")
        if row["icon"] and row["icon"] not in ICON_WHITELIST:
            warnings.append(f"{at}: 图标 {row['icon']} 不在白名单，将回落默认图标")
        if row["navKind"] and row["navKind"] not in NAV_KINDS:
            errors.append(f"{at}: navKind 须为 {'/'.join(NAV_KINDS)}")
        if row["whenCapability"] and row["whenCapability"] not in CAPABILITY_KEYS:
            warnings.append(f"{at}: 能力键 {row['whenCapability']} 不在探测矩阵，条件恒真")
        if row["navKind"] == "menu":
            menu_ids.add(rid)
        if row["default"] is None and row.get("_default_raw"):
            errors.append(f"{at}: default 值无法识别（请填 是/否）")

    for i, row in enumerate(rows):
        at = f"第 {i + 2} 行（{row['id'] or '?'}）"
        if row["navKind"] == "child":
            if not row["navGroup"]:
                errors.append(f"{at}: child 项必须填 navGroup（指向 menu 父项 id）")
            elif row["navGroup"] not in seen:
                errors.append(f"{at}: navGroup={row['navGroup']} 没有对应声明行")
            elif row["navGroup"] not in menu_ids:
                errors.append(f"{at}: navGroup={row['navGroup']} 不是 menu 类型（仅 menu 可嵌套）")
        if row["navGroup"] and row["navKind"] and row["navKind"] != "child":
            warnings.append(f"{at}: navKind={row['navKind']} 时 navGroup 不生效")
        if row["propsJson"]:
            try:
                json.loads(row["propsJson"])
            except (json.JSONDecodeError, TypeError):
                errors.append(f"{at}: propsJson 不是合法 JSON")

    defaults = [r for r in rows if r["default"] is True]
    if len(defaults) > 1:
        warnings.append(f"落地页（default）声明了 {len(defaults)} 个，仅第一个生效")
    return errors, warnings


def to_migration(rows: list[dict[str, Any]]) -> dict[str, Any]:
    errors, _ = validate_rows(rows)
    if errors:
        raise ValueError("注册表校验未通过：\n" + "\n".join(errors))
    views = []
    for row in rows:
        props: dict[str, Any] = {}
        if row["navKind"] and row["navKind"] != "tab":
            nav: dict[str, Any] = {"kind": row["navKind"]}
            if row["navKind"] == "child" and row["navGroup"]:
                nav["group"] = row["navGroup"]
            if row["navKind"] == "header" and row["navOrder"] is not None:
                nav["order"] = row["navOrder"]
            props["nav"] = nav
        if row["propsJson"]:
            props.update(json.loads(row["propsJson"]))
        views.append({
            "id": row["id"], "type": row["type"], "title": row["title"],
            **({"icon": row["icon"]} if row["icon"] else {}),
            **({"default": True} if row["default"] else {}),
            **({"sort": row["sort"]} if row["sort"] is not None else {}),
            **({"when": {"capability": row["whenCapability"]}} if row["whenCapability"] else {}),
            "props": props,
        })
    return {"version": 1, "kind": "registration", "generated_at": _now(), "views": views}


# ── 导入（审核后） ─────────────────────────────────────────

class RegistrationService:
    def __init__(self, site_repo: SiteRepo, db) -> None:
        self._site_repo = site_repo
        self._db = db

    def preview(self, data: bytes) -> dict[str, Any]:
        rows = parse_workbook(data)
        errors, warnings = validate_rows(rows)
        out: dict[str, Any] = {"ok": not errors, "errors": errors, "warnings": warnings, "rows": len(rows)}
        if not errors:
            out["migration"] = to_migration(rows)
        return out

    def apply(self, migration: dict[str, Any]) -> dict[str, Any]:
        if migration.get("kind") != "registration" or migration.get("version") != 1:
            raise ValueError("仅支持 kind=registration 且 version=1 的 migration")
        rows = [
            {
                "id": v.get("id", ""), "type": v.get("type", ""), "title": v.get("title", ""),
                "icon": v.get("icon"), "navKind": (v.get("props", {}).get("nav") or {}).get("kind"),
                "navGroup": (v.get("props", {}).get("nav") or {}).get("group"),
                "navOrder": (v.get("props", {}).get("nav") or {}).get("order"),
                "default": bool(v.get("default")), "sort": v.get("sort"),
                "whenCapability": (v.get("when") or {}).get("capability"),
                "propsJson": json.dumps(
                    {k: val for k, val in v.get("props", {}).items() if k != "nav"}, ensure_ascii=False)
                or None,
            }
            for v in migration.get("views", [])
        ]
        errors, warnings = validate_rows(rows)
        if errors:
            raise ValueError("导入前校验未通过（不信任客户端数据）：\n" + "\n".join(errors))

        applied: list[str] = []
        with self._db.pool.connection() as conn:
            for row, view in zip(rows, migration["views"]):
                self._site_repo.upsert({
                    "id": view["id"], "type": view["type"], "title": view["title"],
                    "icon": view.get("icon"),
                    "when": view.get("when"),
                    "props": view.get("props", {}),
                    "sort": view.get("sort", 100),
                    "default": bool(view.get("default")),
                }, is_builtin=False)
                applied.append(view["id"])
            conn.execute(
                "INSERT INTO registration_migration (source, payload, checksum, applied)"
                " VALUES (%s, %s, %s, TRUE)",
                ("upload", json.dumps(migration, ensure_ascii=False), _checksum(migration["views"])),
            )
        return {"applied": applied, "warnings": warnings}
