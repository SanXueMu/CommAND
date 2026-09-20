"""AO4：视图导出 xlsx 的数值化——年份/月份/金额等纯数字串写数值单元格。

Excel 对文本单元格按字典序排序（「10」排在「9」前），故数值化是月份/年份可正确
排序的前提；同时超长数字与前导零串保留文本，避免精度丢失与丢位。
"""
from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def tool():
    spec = importlib.util.spec_from_file_location(
        "_ao4_export_xlsx", ROOT / "tools" / "records" / "export_xlsx" / "main.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_pure_number_strings_become_numeric_cells(tmp_path, monkeypatch, tool):
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path))
    columns = ["年份", "月份", "借方金额", "凭证号", "附件", "摘要", "流水号", "编号"]
    rows = [
        {"年份": "2002", "月份": "9", "借方金额": "511.84", "凭证号": "总1号",
         "附件": "1张", "摘要": "购文具", "流水号": "12345678901234567890", "编号": "007"},
        {"年份": "2002", "月份": "10", "借方金额": "0.42", "凭证号": "总2号",
         "附件": "", "摘要": "现金", "流水号": "1", "编号": "12"},
    ]
    out = tool.run({"rows": rows, "columns": columns, "name": "数值化", "sheet_name": "视图"},
                   None, lambda e: None)
    ws = load_workbook(out["file"]).active
    assert ws.cell(2, 1).value == 2002 and isinstance(ws.cell(2, 1).value, int)
    assert ws.cell(2, 2).value == 9 and not isinstance(ws.cell(2, 2).value, str)
    assert ws.cell(3, 2).value == 10, "两位数月份也须为数值（文本序 10 会排在 9 前）"
    assert ws.cell(2, 3).value == 511.84
    assert ws.cell(2, 4).value == "总1号"
    assert ws.cell(2, 5).value == "1张"
    assert ws.cell(2, 6).value == "购文具"
    assert ws.cell(2, 7).value == "12345678901234567890", "超长数字保留文本防精度丢失"
    assert ws.cell(2, 8).value == "007", "前导零串保留文本不丢位"
    assert ws.cell(3, 8).value == 12


@pytest.mark.parametrize("given,expected", [
    ("2003年1月记账凭证（2）", "2003年1月记账凭证（2）.xlsx"),
    ("2003年1月记账凭证（2）.xlsx", "2003年1月记账凭证（2）.xlsx"),
    ("合并导出12库.xls", "合并导出12库.xlsx"),
    ("", "视图导出.xlsx"),
])
def test_export_name_never_doubles_suffix(tmp_path, monkeypatch, tool, given, expected):
    """P2：调用方带不带 .xlsx/.xls 后缀，产物名都只出现一次后缀。"""
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path))
    out = tool.run({"rows": [{"a": "1"}], "columns": ["a"], "name": given, "sheet_name": "视图"},
                   None, lambda e: None)
    assert Path(out["file"]).name == expected
