"""工具域测试：10 个 v2 工具的纯函数链 + 假客户端翻译引擎 + manifest 协议校验。

约定：ctx 用 SimpleNamespace 伪造（handle/keys）；COMMAND_DATA_DIR 指 tmp_path；
LLM 不出网——make_client 被替换为返回 FakeClient（按 prompt 内容应答）。
"""
from __future__ import annotations

import importlib.util
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest

REPO = Path(__file__).resolve().parent.parent

TOOLS = [
    "tools/xlsx/extract_values",
    "tools/pdf/extract_pages",
    "tools/txt/extract_text",
    "tools/ocrdb/extract_units",
    "tools/table/classify_columns",
    "tools/text/dedup_values",
    "tools/text/llm_translate",
    "tools/text/verify_fidelity",
    "tools/xlsx/backfill_dict",
    "tools/docx/render_bilingual",
]


def load_tool(rel: str):
    path = REPO / rel / "main.py"
    spec = importlib.util.spec_from_file_location("tool_" + rel.replace("/", "_"), path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def ctx(monkeypatch, tmp_path):
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path))
    return SimpleNamespace(
        handle="test-handle",
        keys={"main": {"provider": "dashscope", "base_url": "https://x/v1",
                       "api_key": "sk-test", "is_default": True}},
    )


class _Msg:
    def __init__(self, content):
        self.content = content


class _Choice:
    def __init__(self, content):
        self.message = _Msg(content)


class _Resp:
    def __init__(self, content):
        self.choices = [_Choice(content)]
        self.usage = None


class FakeClient:
    """按 prompt 应答：分隔符批翻协议 / JSON 兜底 / 重译修正。"""

    def __init__(self, retranslator=None):
        self.calls: list[dict] = []
        self.retranslator = retranslator

    _SEP_RE = re.compile(r"###(\d+)###\n?(.*?)(?=###\d+###|$)", re.S)

    @staticmethod
    def _segments_from_prompt(prompt: str) -> list[str]:
        found = {int(m.group(1)): m.group(2).strip() for m in FakeClient._SEP_RE.finditer(prompt)}
        return [found[k] for k in sorted(found)]

    def _make(self, responder):
        outer = self

        class Completions:
            def create(self, *, model=None, messages=None, **kw):
                outer.calls.append({"model": model, "messages": messages})
                return _Resp(responder(model, messages[-1]["content"]))

        class Chat:
            completions = Completions()

        return Chat()

    def __post_init__(self):
        raise NotImplementedError


def separator_client(target="中"):
    fake = FakeClient.__new__(FakeClient)
    fake.calls = []
    fake.retranslator = None

    def respond(model, prompt):
        segs = FakeClient._segments_from_prompt(prompt)
        out = "\n".join(f"###{i + 1}###\n{s}{target}" for i, s in enumerate(segs))
        return out or "###1###\nx"

    fake.chat = fake._make(respond)
    return fake


def retranslate_client(bad_first=1):
    """首 n 次给缺数字的坏译文，之后给合格译文（测质检重译环）。"""
    fake = FakeClient.__new__(FakeClient)
    fake.calls = []
    state = {"n": bad_first}
    fake.retranslator = state

    def respond(model, prompt):
        if "The previous translation" in prompt or "Source:" in prompt:
            state["n"] -= 1
            if "###" in prompt:  # 重译仍走批翻分隔符格式
                return "###1###\n净利润 12.5 百万元（重译合格）"
            return "净利润 12.5 百万元（重译合格）"
        segs = FakeClient._segments_from_prompt(prompt)
        parts = []
        for i, s in enumerate(segs):
            parts.append(f"###{i + 1}###\nNet profit million yuan" if state["n"] > 0 else f"###{i + 1}###\nNet profit 12.5 million yuan")
        return "\n".join(parts)

    fake.chat = fake._make(respond)
    return fake


# ---------- manifest 协议校验 ----------


def test_all_manifests_valid():
    import tomllib

    from core.protocol import ToolManifest

    for rel in TOOLS:
        raw = tomllib.loads((REPO / rel / "tool.toml").read_text(encoding="utf-8"))
        assert (REPO / rel / raw["io"]["input_schema"]).exists(), rel
        assert (REPO / rel / raw["io"]["output_schema"]).exists(), rel
        manifest = ToolManifest.from_toml(REPO / rel / "tool.toml")
        assert manifest.io.input_schema and manifest.io.output_schema  # 解析为 dict
        assert manifest.runtime.kind == "inproc"


# ---------- 提取域 ----------


def test_xlsx_extract_classify_dedup_chain(ctx, tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "底稿"
    ws.append(["项目", "Amount", "Date", "说明"])
    ws.append(["营业收入", "100.00", "2023-11-16", "Service fee"])
    ws.append(["营业成本", "80.00", "2023-11-17", "Consulting fee"])
    wb.save(p := tmp_path / "t.xlsx")

    extracted = load_tool("tools/xlsx/extract_values").run({"file": str(p)}, ctx, lambda e: None)
    assert extracted["file_hash"]
    unit = extracted["units"][0]
    assert unit["unit_type"] == "table" and unit["rows"][0][0] == "项目"

    cls = load_tool("tools/table/classify_columns").run(
        {"units": extracted["units"]}, ctx, lambda e: None
    )
    classes = cls["col_classes"][unit["unit_id"]]
    by_col = {c["col"]: c["class"] for c in classes}
    assert by_col[1] == "skip" and by_col[2] == "skip"      # 纯数字列 / 日期列不译
    assert by_col[0] == "enum" and by_col[3] == "enum"      # 低基数文本列（命中字典缓存复用）
    assert cls["segments"] == [
        "项目", "Amount", "Date", "说明",                    # 表头（含字母即收割）
        "营业收入", "Service fee", "营业成本", "Consulting fee",  # 数据：仅非 skip 列
    ]

    dedup = load_tool("tools/text/dedup_values").run({"segments": cls["segments"]}, ctx, lambda e: None)
    assert dedup["unique"] == cls["segments"]
    assert dedup["index_map"] == list(range(8))


def test_dedup_dates_placeholder(ctx):
    mod = load_tool("tools/text/dedup_values")
    out = mod.run(
        {"segments": ["2023-11-16 报告", "2023-11-17 报告", "纯文本"]}, ctx, lambda e: None
    )
    assert out["unique"] == ["[[DATE_1]] 报告", "纯文本"]
    assert out["index_map"] == [0, 0, 1]
    assert out["date_maps"][0] == {"[[DATE_1]]": "2023-11-16"}
    assert out["date_maps"][1] == {"[[DATE_1]]": "2023-11-17"}
    assert out["date_maps"][2] == {}


def test_pdf_and_txt_extract(ctx, tmp_path):
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "审计报告")
    doc.save(pdf := tmp_path / "a.pdf")
    doc.close()

    units = load_tool("tools/pdf/extract_pages").run({"file": str(pdf)}, ctx, lambda e: None)["units"]
    assert units[0]["meta"]["page"] == 1

    (txt := tmp_path / "b.txt").write_text("hello world", encoding="utf-8")
    txt_out = load_tool("tools/txt/extract_text").run({"file": str(txt)}, ctx, lambda e: None)
    assert txt_out["units"][0]["text"] == "hello world"


def test_ocrdb_extract(ctx, tmp_path):
    import json

    from command_shared import ocr_storage

    long_text = "长" * 250
    # AJ1 后查询含 seq 列——桩库必须走 ocr_storage 建表（自动迁移不认外来同名表）
    p = tmp_path / "ocr.sqlite"
    db = ocr_storage.connect(p)
    ocr_storage.initialize(db)
    db.execute(
        "INSERT OR REPLACE INTO records (file_hash, source_path, row_number, seq, page_number, data) "
        "VALUES ('h1', 'scan.pdf', 1, 0, 1, ?)",
        (json.dumps({"发票号": "INV-1", "金额": "100.00", "备注": long_text}),),
    )
    db.commit()
    db.close()
    units = load_tool("tools/ocrdb/extract_units").run({"file": str(p)}, ctx, lambda e: None)["units"]
    table_units = [u for u in units if u["unit_type"] == "table"]
    text_units = [u for u in units if u["unit_type"] == "text"]
    assert table_units and table_units[0]["rows"][2] == ["发票号 (p1)", "INV-1"]
    assert text_units and text_units[0]["text"] == long_text  # 长值分流


# ---------- 翻译引擎（假客户端） ----------


def test_translate_identity_and_separator(ctx, monkeypatch):
    mod = load_tool("tools/text/llm_translate")
    fake = separator_client()
    monkeypatch.setattr(mod, "make_client", lambda key, base: fake)
    out = mod.run(
        {"segments": ["INV-2024-001", "hello world", "", "fee 100"], "key_name": "main",
         "use_cache": False},
        ctx, lambda e: None,
    )
    assert out["translations"][0] == "INV-2024-001"  # 编号直存
    assert out["translations"][1] == "hello world中"
    assert out["translations"][2] == ""  # 空段直通
    assert out["statuses"] == ["ok", "ok", "ok", "ok"]
    assert out["cache_hits"] == 0


def test_translate_review_and_retranslate(ctx, monkeypatch):
    mod = load_tool("tools/text/llm_translate")
    fake = retranslate_client(bad_first=1)
    monkeypatch.setattr(mod, "make_client", lambda key, base: fake)
    out = mod.run(
        {"segments": ["Net profit was 12.5 million yuan"], "use_cache": False}, ctx, lambda e: None
    )
    assert out["statuses"][0] == "ok"
    assert "12.5" in out["translations"][0]
    assert len(fake.calls) >= 2  # 批翻(坏) → 重译(合格)


def test_translate_requires_key(ctx):
    import copy

    from core.errors import ToolDomainError

    mod = load_tool("tools/text/llm_translate")
    empty = SimpleNamespace(handle="t", keys={})
    with pytest.raises(ToolDomainError):
        mod.run({"segments": ["x"], "use_cache": False}, empty, lambda e: None)


def test_translate_cache_roundtrip(ctx, monkeypatch):
    mod = load_tool("tools/text/llm_translate")
    fake = separator_client()
    monkeypatch.setattr(mod, "make_client", lambda key, base: fake)
    mod.run({"segments": ["hello world"], "use_cache": True}, ctx, lambda e: None)
    n_calls = len(fake.calls)
    out2 = mod.run({"segments": ["hello world"], "use_cache": True}, ctx, lambda e: None)
    assert out2["cache_hits"] == 1 and len(fake.calls) == n_calls  # 二跑零 API


# ---------- 质检 / 回填 / 渲染 ----------


def test_verify_tool(ctx):
    out = load_tool("tools/text/verify_fidelity").run(
        {"sources": ["fee 100", "fee 200"], "translations": ["费用 100", "费用"]},
        ctx, lambda e: None,
    )
    assert out["statuses"] == ["ok", "review"]
    assert "missing_number" in out["reasons"][1]


def _chain_fixtures():
    units = [{"unit_id": "s1", "unit_type": "table",
              "rows": [["项目", "Amount"], ["营业收入", "100.00"], ["管理费用", "20.00"]]}]
    cls = load_tool("tools/table/classify_columns").run({"units": units}, None, lambda e: None)
    dedup = load_tool("tools/text/dedup_values").run({"segments": cls["segments"]}, None, lambda e: None)
    return units, cls, dedup


def test_xlsx_backfill_join(ctx, tmp_path):
    """dict 版式（旧行为）：每单元一张 sheet + 末尾对照字典。"""
    units, cls, dedup = _chain_fixtures()
    n = len(dedup["unique"])
    backfill = load_tool("tools/xlsx/backfill_dict")
    out = backfill.run(
        {"units": units, "col_classes": cls["col_classes"], "ranges": cls["ranges"],
         "segments": cls["segments"], "index_map": dedup["index_map"],
         "date_maps": dedup["date_maps"],
         "translations": [f"译{i}" for i in range(n)],
         "statuses": ["ok"] * (n - 1) + ["review"], "file": "底稿.xlsx",
         "layout": "dict"},
        ctx, lambda e: None,
    )
    from openpyxl import load_workbook

    assert out["layout"] == "dict"
    wb = load_workbook(out["path"])
    ws = wb["s1"]
    assert ws.cell(1, 1).value == "译0"   # 表头 项目 → ok
    assert ws.cell(1, 2).value == "译1"   # 表头 Amount → ok
    assert ws.cell(2, 1).value == "译2"   # 营业收入 → ok
    assert ws.cell(3, 1).value.startswith("⚠️")  # review 格保留原文
    assert ws.cell(3, 1).value == "⚠️ 管理费用"
    assert wb.sheetnames[-1] == "对照字典"


def test_xlsx_backfill_sheets_layout(ctx, tmp_path):
    """sheets 版式（默认）：原 sheet 之后插入 <sheet名>_翻译结果，原文不动，无对照字典。"""
    from openpyxl import Workbook, load_workbook

    src = tmp_path / "投标资料.xlsx"
    wb0 = Workbook()
    ws1 = wb0.active
    ws1.title = "Sheet1"
    ws1.append(["项目", "金额"])
    ws1.append(["营业收入", "100.00"])
    ws1.append(["管理费用", "20.00"])
    ws2 = wb0.create_sheet("Sheet2")
    ws2.append(["条款", "说明"])
    ws2.append(["Payment", "within 30 days"])
    wb0.save(src)

    units = [
        {"unit_id": "Sheet1", "unit_type": "table",
         "rows": [["项目", "金额"], ["营业收入", "100.00"], ["管理费用", "20.00"]]},
        {"unit_id": "Sheet2", "unit_type": "table",
         "rows": [["条款", "说明"], ["Payment", "within 30 days"]]},
    ]
    cls = load_tool("tools/table/classify_columns").run({"units": units}, None, lambda e: None)
    dedup = load_tool("tools/text/dedup_values").run({"segments": cls["segments"]}, None, lambda e: None)
    n = len(dedup["unique"])
    out = load_tool("tools/xlsx/backfill_dict").run(
        {"units": units, "col_classes": cls["col_classes"], "ranges": cls["ranges"],
         "segments": cls["segments"], "index_map": dedup["index_map"], "date_maps": dedup["date_maps"],
         "translations": [f"译文{i}" for i in range(n)],
         "statuses": ["ok"] * n, "file": str(src)},
        ctx, lambda e: None,
    )

    assert out["layout"] == "sheets"
    assert out["name"] == "投标资料_中文.xlsx"
    assert out["sheets_made"] == 2
    wb = load_workbook(out["path"])
    # 每个原 sheet 之后紧跟其译文 sheet；无对照字典
    assert wb.sheetnames == ["Sheet1", "Sheet1_翻译结果", "Sheet2", "Sheet2_翻译结果"]
    # 译文按原坐标回填（期望值由 dedup 的唯一段序推出，不写死下标）
    idx = {s: i for i, s in enumerate(dedup["unique"])}
    w1 = wb["Sheet1_翻译结果"]
    assert w1.cell(1, 1).value == f"译文{idx['项目']}"
    assert w1.cell(2, 1).value == f"译文{idx['营业收入']}"
    assert w1.cell(3, 1).value == f"译文{idx['管理费用']}"
    w2 = wb["Sheet2_翻译结果"]
    assert w2.cell(1, 1).value == f"译文{idx['条款']}"
    assert w2.cell(2, 1).value == f"译文{idx['Payment']}"
    assert w2.cell(2, 2).value == f"译文{idx['within 30 days']}"
    # 原 sheet 一字不动
    o1 = wb["Sheet1"]
    assert o1.cell(1, 1).value == "项目"
    assert o1.cell(2, 1).value == "营业收入"
    # 源文件未被修改（仍只有两个 sheet）
    assert load_workbook(src).sheetnames == ["Sheet1", "Sheet2"]



def test_docx_render(ctx, tmp_path):
    units = [{"unit_id": "p1", "unit_type": "text", "text": "Audit Summary",
              "meta": {"page": 1, "kind": "text"}}]
    cls = load_tool("tools/table/classify_columns").run({"units": units}, None, lambda e: None)
    dedup = load_tool("tools/text/dedup_values").run({"segments": cls["segments"]}, None, lambda e: None)
    render = load_tool("tools/docx/render_bilingual")
    out = render.run(
        {"units": units, "ranges": cls["ranges"], "index_map": dedup["index_map"],
         "date_maps": dedup["date_maps"], "translations": ["审计总结"],
         "statuses": ["ok"], "file": "报告.pdf"},
        ctx, lambda e: None,
    )
    from docx import Document

    doc = Document(out["path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "审计总结" in text and "第 1 页" in text


def test_docx_render_table_unit(ctx):
    units = [{"unit_id": "t1", "unit_type": "table", "rows": [["项目", "Amount"], ["营业收入", "100.00"]]}]
    cls = load_tool("tools/table/classify_columns").run({"units": units}, None, lambda e: None)
    dedup = load_tool("tools/text/dedup_values").run({"segments": cls["segments"]}, None, lambda e: None)
    n = len(dedup["unique"])
    out = load_tool("tools/docx/render_bilingual").run(
        {"units": units, "ranges": cls["ranges"], "index_map": dedup["index_map"],
         "date_maps": dedup["date_maps"], "col_classes": cls["col_classes"],
         "translations": [f"译{i}" for i in range(n)], "statuses": ["ok"] * n},
        ctx, lambda e: None,
    )
    from docx import Document

    doc = Document(out["path"])
    assert doc.tables and doc.tables[0].cell(1, 0).text == "译2"


def test_ocrdb_extract_records_merge_multi_db(ctx, tmp_path):
    """AJ1：file 传数组即多库合并导出——逐库读取、来源分组、同页按 seq 排序。"""
    import json
    import sqlite3

    from command_shared import ocr_storage

    def _mk_db(name: str, rows: list[tuple[str, int, int, dict]]) -> str:
        p = tmp_path / name
        conn = ocr_storage.connect(p)
        ocr_storage.initialize(conn)
        for src, page, seq, data in rows:
            conn.execute(
                "INSERT OR REPLACE INTO records (file_hash, source_path, row_number, seq, page_number, data) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (f"h-{src}", src, page, seq, page, json.dumps(data, ensure_ascii=False)),
            )
        conn.commit()
        conn.close()
        return str(p)

    a = _mk_db("a.ocr_results.db", [
        ("/2002年3月记账凭证_1_.pdf", 1, 0, {"金额": "100"}),
        ("/2002年3月记账凭证_1_.pdf", 1, 1, {"金额": "200"}),
        ("/2002年3月记账凭证_2_.pdf", 2, 0, {"金额": "300"}),
    ])
    b = _mk_db("b.ocr_results.db", [("/2002年4月记账凭证_1_.pdf", 1, 0, {"金额": "999"})])

    out = load_tool("tools/ocrdb/extract_units").run(
        {"file": [a, b], "mode": "records"}, ctx, lambda e: None)
    assert out["file"] == "合并导出(2库)"
    amounts = [r["金额"] for r in out["records"]]
    # 同页按 seq（100→200），跨库来源分组有序（a 全部在前、b 在后）
    assert amounts == ["100", "200", "300", "999"]
    assert out["records"][0]["来源文件"] == "2002年3月记账凭证_1_.pdf"
    assert out["records"][-1]["来源文件"] == "2002年4月记账凭证_1_.pdf"
    # AO：年份/月份拆两列纯数字串（来源仍是文件名；Excel 数值排序正确）
    assert out["records"][0]["年份"] == "2002" and out["records"][0]["月份"] == "03"
    assert out["records"][2]["年份"] == "2002" and out["records"][2]["月份"] == "03"
    assert out["records"][-1]["年份"] == "2002" and out["records"][-1]["月份"] == "04"


def test_ocrdb_extract_records_month_order(ctx, tmp_path):
    """AO2：多库合并行序按（年份, 月份）数字序——10 月不得排在 2 月之前。"""
    import json

    from command_shared import ocr_storage

    def _mk_db(name: str, src: str) -> str:
        p = tmp_path / name
        conn = ocr_storage.connect(p)
        ocr_storage.initialize(conn)
        conn.execute(
            "INSERT OR REPLACE INTO records (file_hash, source_path, row_number, seq, page_number, data) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (f"h-{src}", src, 1, 0, 1, json.dumps({"金额": "1"}, ensure_ascii=False)))
        conn.commit()
        conn.close()
        return str(p)

    # 目录名（含 2002年10月 / 2002年2月）按字典序恰好是「10 月在 2 月前」
    oct_db = _mk_db("oct.ocr_results.db", "/2002年10月记账凭证_1_.pdf")
    feb_db = _mk_db("feb.ocr_results.db", "/2002年2月记账凭证_1_.pdf")
    out = load_tool("tools/ocrdb/extract_units").run(
        {"file": [oct_db, feb_db], "mode": "records"}, ctx, lambda e: None)
    months = [r["月份"] for r in out["records"]]
    assert months == ["02", "10"], f"行序应按月份数字序: {months}"
    assert [r["年份"] for r in out["records"]] == ["2002", "2002"]
    # 无年月模式的来源：两列留空且排最后
    other = _mk_db("other.ocr_results.db", "/无年月标记.pdf")
    out2 = load_tool("tools/ocrdb/extract_units").run(
        {"file": [other, feb_db], "mode": "records"}, ctx, lambda e: None)
    assert out2["records"][-1]["年份"] == "" and out2["records"][-1]["月份"] == ""


def test_ocrdb_extract_records_range_month_disambiguation(ctx, tmp_path):
    """AR2：区间文件名（1993年9月-10月）月份列直接写补零范围串——"09-10"。

    凭证日期识别率不足，逐条消歧不可靠；文件名范围即权威口径。
    所有行同值；排序键取区间首月（同文件聚组、页序稳定）；_sk 不进导出产物。
    """
    import json

    from command_shared import ocr_storage

    def _mk_db(name: str, rows: list[tuple[str, int, dict]]) -> str:
        p = tmp_path / name
        conn = ocr_storage.connect(p)
        ocr_storage.initialize(conn)
        for src, page, data in rows:
            conn.execute(
                "INSERT OR REPLACE INTO records (file_hash, source_path, row_number, seq, page_number, data) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (f"h-{src}-{page}", src, page, 0, page, json.dumps(data, ensure_ascii=False)))
        conn.commit()
        conn.close()
        return str(p)

    rng = "/蜀棱公司1993年9月-10月记账凭证.pdf"
    db = _mk_db("r.ocr_results.db", [
        (rng, 1, {"日期": "1993年9月15日", "金额": "a"}),
        (rng, 2, {"日期": "1993年1月5日", "金额": "misread"}),
        (rng, 3, {"金额": "nodate"}),
        (rng, 4, {"日期": "1993-10-5", "金额": "b"}),
    ])
    out = load_tool("tools/ocrdb/extract_units").run(
        {"file": db, "mode": "records"}, ctx, lambda e: None)
    # 区间文件：全部行月份=范围串（日期字段不参与判定）
    for r in out["records"]:
        assert r["月份"] == "09-10" and r["年份"] == "1993"
    # 同文件聚组且保持页序
    order = [r["金额"] for r in out["records"]]
    assert order == ["a", "misread", "nodate", "b"], order
    assert all("_sk" not in r for r in out["records"])


def test_ocrdb_extract_records_cross_year_range(ctx, tmp_path):
    """AR2：跨年区间（1993年11月-1994年1月）——月份列写 "11-01"，年份取文件名年。"""
    import json

    from command_shared import ocr_storage

    p = tmp_path / "x.ocr_results.db"
    conn = ocr_storage.connect(p)
    ocr_storage.initialize(conn)
    src = "/1993年11月-1994年1月装订册.pdf"
    conn.execute(
        "INSERT OR REPLACE INTO records (file_hash, source_path, row_number, seq, page_number, data) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (f"h-{src}", src, 1, 0, 1, json.dumps({"日期": "1994年1月8日"}, ensure_ascii=False)))
    conn.commit()
    conn.close()
    out = load_tool("tools/ocrdb/extract_units").run(
        {"file": str(p), "mode": "records"}, ctx, lambda e: None)
    assert out["records"][0]["年份"] == "1993" and out["records"][0]["月份"] == "11-01"


def test_ocrdb_extract_records_single_month_filename_authoritative(ctx, tmp_path):
    """AR 前后行为不变式：单月文件名下文件名权威，凭证日期不参与月份判定。"""
    import json

    from command_shared import ocr_storage

    p = tmp_path / "s.ocr_results.db"
    conn = ocr_storage.connect(p)
    ocr_storage.initialize(conn)
    src = "/2002年3月记账凭证_1_.pdf"
    conn.execute(
        "INSERT OR REPLACE INTO records (file_hash, source_path, row_number, seq, page_number, data) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (f"h-{src}", src, 1, 0, 1, json.dumps({"日期": "2002年5月1日"}, ensure_ascii=False)))
    conn.commit()
    conn.close()
    out = load_tool("tools/ocrdb/extract_units").run(
        {"file": str(p), "mode": "records"}, ctx, lambda e: None)
    assert out["records"][0]["年份"] == "2002" and out["records"][0]["月份"] == "03"
