"""docx.render.bilingual 回归：空单元不再越界崩溃 + col_classes 正确消费（表格逐格译文）。

线上事故（2026-09-14，PDF 版式/文档流第 6 步 `IndexError: list index out of range`）：
`pdf.extract.pages` 对**每一页**都生成 text 单元（空页也生成，标 no_text），而
`harvest_units` 对空文本给 `count=0`；渲染端原先不看 count，直接 `index_map[rng["start"]]`
—— 空页落在末尾时 start 恰好等于 len(index_map) → 越界崩溃。
同时该工具漏接线 col_classes（xlsx / docx.translated 两个兄弟工具都传了），
导致双语 docx 的表格只有表头被翻、数据格回落原文。
"""
from __future__ import annotations

import importlib
import types
from pathlib import Path


def _tool():
    return importlib.import_module("tools.docx.render_bilingual.main")


def _run(payload: dict, tmp_path: Path, handle: str = "h") -> dict:
    return _tool().run({"file": "src.pdf", **payload},
                       types.SimpleNamespace(handle=handle), lambda e: None)


def _texts(path: Path) -> list[str]:
    from docx import Document

    document = Document(str(path))
    out = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            out.extend(cell.text for cell in row.cells)
    return out


def test_trailing_empty_text_unit_does_not_crash(tmp_path, monkeypatch):
    """空页 text 单元（count=0）落在末尾时，曾致 index_map[len] 越界。"""
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    out = _run({
        "units": [
            {"unit_id": "1", "unit_type": "text", "text": "Hello", "meta": {"page": 1}},
            {"unit_id": "2", "unit_type": "text", "text": "",
             "meta": {"page": 2, "no_text": True}},
        ],
        "ranges": [
            {"unit_id": "1", "start": 0, "count": 1},
            {"unit_id": "2", "start": 1, "count": 0},
        ],
        "index_map": [0],
        "date_maps": [{}],
        "translations": ["你好"],
        "statuses": ["ok"],
    }, tmp_path)
    assert Path(out["path"]).is_file()
    assert "你好" in _texts(Path(out["path"]))


def test_empty_text_unit_in_middle_is_skipped(tmp_path, monkeypatch):
    """中间的空白单元不得消费下一条单元的译文（否则错位串译）。"""
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    out = _run({
        "units": [
            {"unit_id": "1", "unit_type": "text", "text": "First", "meta": {"page": 1}},
            {"unit_id": "2", "unit_type": "text", "text": "",
             "meta": {"page": 2, "no_text": True}},
            {"unit_id": "3", "unit_type": "text", "text": "Second", "meta": {"page": 3}},
        ],
        "ranges": [
            {"unit_id": "1", "start": 0, "count": 1},
            {"unit_id": "2", "start": 1, "count": 0},
            {"unit_id": "3", "start": 1, "count": 1},
        ],
        "index_map": [0, 1],
        "date_maps": [{}, {}],
        "translations": ["第一", "第二"],
        "statuses": ["ok", "ok"],
    }, tmp_path)
    texts = _texts(Path(out["path"]))
    assert "第一" in texts and "第二" in texts


def test_table_cells_use_translations_when_col_classes_given(tmp_path, monkeypatch):
    """col_classes 未接线时表格数据格回落原文（线上双语 docx 表格没译文）。"""
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    out = _run({
        "units": [
            {"unit_id": "1T1", "unit_type": "table",
             "rows": [["A", "B"], ["x1", "Alpha"]], "meta": {"page": 1}},
        ],
        "ranges": [{"unit_id": "1T1", "start": 0, "count": 3}],
        "col_classes": {"1T1": [{"col": 0, "class": "skip"}, {"col": 1, "class": "text"}]},
        "index_map": [0, 1, 2],
        "date_maps": [{}, {}, {}],
        "translations": ["甲", "乙", "阿尔法"],
        "statuses": ["ok", "ok", "ok"],
    }, tmp_path)
    cells = _texts(Path(out["path"]))
    assert "阿尔法" in cells, "表格数据格应使用译文"
    assert "Alpha" not in cells, "数据格不应回落原文"


def test_table_without_col_classes_still_renders(tmp_path, monkeypatch):
    """缺 col_classes 时不得越界：按 count 收敛，未映射的格保留原文。"""
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    out = _run({
        "units": [
            {"unit_id": "1T1", "unit_type": "table",
             "rows": [["A", "B"], ["x1", "Alpha"]], "meta": {"page": 1}},
        ],
        "ranges": [{"unit_id": "1T1", "start": 0, "count": 3}],
        "index_map": [0],
        "date_maps": [{}],
        "translations": ["甲"],
        "statuses": ["ok"],
    }, tmp_path)
    cells = _texts(Path(out["path"]))
    assert "Alpha" in cells


def test_missing_statuses_and_date_maps_do_not_crash(tmp_path, monkeypatch):
    """statuses / date_maps 比 index_map 短（历史 run 缺键）时回退默认值而非崩溃。"""
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    out = _run({
        "units": [
            {"unit_id": "1", "unit_type": "text", "text": "Hello", "meta": {"page": 1}},
        ],
        "ranges": [{"unit_id": "1", "start": 0, "count": 1}],
        "index_map": [0],
        "date_maps": [],
        "translations": ["你好"],
        "statuses": [],
    }, tmp_path)
    assert "你好" in _texts(Path(out["path"]))
