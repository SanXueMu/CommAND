"""docx.render.translated / command_shared.docx_render 测试（overlay 与 bilingual 双模式）。"""
from __future__ import annotations

import importlib
import re
import types
from pathlib import Path

import pytest

from command_shared.docx_units import extract_units
from command_shared.table import DEFAULT_SKIP_RE, classify_columns, harvest_units

SKIP_RE = re.compile(DEFAULT_SKIP_RE)


def _tool():
    return importlib.import_module("tools.docx.render_translated.main")


def _tiny_png(path: Path) -> Path:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=20, height=20)
    page.insert_text((4, 12), "S", fontsize=8)
    page.get_pixmap(dpi=40).save(str(path))
    doc.close()
    return path


def _field_paragraph(document, label: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = document.add_paragraph()
    paragraph.add_run(label)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    return paragraph


def _source_docx(path: Path, tmp_path: Path) -> Path:
    """p0 标题 / p1 正文 / p2 空段 / p3 纯数字 / p4 域段 / t5 表格 / p6 图文段 / p7 尾段。"""
    from docx import Document

    doc = Document()
    doc.add_heading("Audit Engagement Letter", level=1)
    doc.add_paragraph("This agreement is made between ZTH CPA and the Client.")
    doc.add_paragraph("")
    doc.add_paragraph("2026-01-01")
    _field_paragraph(doc, "Page ")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Audit fee"
    table.cell(1, 1).text = "123456"
    image_para = doc.add_paragraph()
    image_para.add_run("Seal image: ")
    image_para.add_run().add_picture(str(_tiny_png(tmp_path / "seal.png")))
    doc.add_paragraph("Signature: Zhang Yi")
    doc.save(path)
    return path


def _chain(src: Path) -> dict:
    """走真实链：extract → classify → harvest → 造译文（index_map/date_maps 置 None）。"""
    from docx import Document

    units = extract_units(Document(str(src)))["units"]
    col_classes = {u["unit_id"]: classify_columns(u["rows"], SKIP_RE)
                   for u in units if u["unit_type"] == "table"}
    harvested = harvest_units(units, col_classes)
    segments = harvested["segments"]
    return {
        "file": str(src), "units": units, "ranges": harvested["ranges"],
        "col_classes": col_classes, "index_map": None, "date_maps": None,
        "translations": [f"【译】{s}" for s in segments],
        "statuses": ["ok"] * len(segments),
    }


def _run(payload: dict, handle: str = "h"):
    return _tool().run(payload, types.SimpleNamespace(handle=handle), lambda e: None)


def _texts(document) -> list[str]:
    return [p.text for p in document.paragraphs]


# ── overlay 原位覆盖 ─────────────────────────────────────────────────────────
def test_overlay_replaces_text_table_and_keeps_structure(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    src = _source_docx(tmp_path / "s.docx", tmp_path)
    out = _run({**_chain(src), "mode": "overlay"})

    assert out["mode"] == "overlay" and Path(out["path"]).is_file()
    assert out["name"].endswith("_原位译文.docx") and out["review"] == 0

    document = Document(out["path"])
    texts = _texts(document)
    assert texts[0] == "【译】Audit Engagement Letter"
    assert document.paragraphs[0].style.name.startswith("Heading")  # 标题级别保留
    assert texts[1].startswith("【译】This agreement")
    assert texts[2] == ""                    # 空段原样
    assert texts[3] == "2026-01-01"          # 纯数字段原样
    assert texts[4].startswith("Page ")      # 域段原样（域元素保留）
    assert document.paragraphs[4]._p.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldSimple")
    assert texts[5] == "【译】Seal image:"    # 图文段：文字替换、图片 run 保留（下测细验）
    assert texts[6] == "【译】Signature: Zhang Yi"

    table = document.tables[0]
    assert table.cell(0, 0).text == "【译】Item"
    assert table.cell(0, 1).text == "【译】Amount"
    assert table.cell(1, 0).text == "【译】Audit fee"
    assert table.cell(1, 1).text == "123456"  # 纯数字列 skip，原样

    # 源文件只读未被改
    assert _texts(Document(str(src)))[0] == "Audit Engagement Letter"


def test_overlay_keeps_image_run(tmp_path, monkeypatch):
    from docx.oxml.ns import qn

    from docx import Document

    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    src = _source_docx(tmp_path / "img.docx", tmp_path)
    out = _run({**_chain(src), "mode": "overlay"})
    document = Document(out["path"])
    image_paragraph = document.paragraphs[5]  # document.paragraphs 不含表格内段落
    assert image_paragraph.text == "【译】Seal image:"
    assert image_paragraph._p.findall(f".//{qn('w:drawing')}"), "图片 run 必须保留"


# ── bilingual 段落对照 ───────────────────────────────────────────────────────
def test_bilingual_keeps_original_and_inserts_after(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    src = _source_docx(tmp_path / "b.docx", tmp_path)
    out = _run({**_chain(src), "mode": "bilingual"})

    assert out["inserted"] >= 5 and out["replaced"] == 0
    document = Document(out["path"])
    texts = _texts(document)
    assert texts[0] == "Audit Engagement Letter"          # 原文保留
    assert texts[1] == "【译】Audit Engagement Letter"     # 紧随其后
    assert document.paragraphs[1].style.name.startswith("Heading")  # 复用标题样式
    run = document.paragraphs[1].runs[0]
    assert run.font.color is not None and str(run.font.color.rgb) == "1F4E79"
    assert texts[2].startswith("This agreement")          # 原正文
    assert texts[3].startswith("【译】This agreement")
    # 表格：原文保留 + 格内新增译文段
    table = document.tables[0]
    assert table.cell(1, 0).paragraphs[0].text == "Audit fee"
    assert table.cell(1, 0).paragraphs[1].text == "【译】Audit fee"
    assert table.cell(1, 1).text == "123456"              # skip 列不动


def test_default_mode_is_bilingual(tmp_path, monkeypatch):
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    src = _source_docx(tmp_path / "d.docx", tmp_path)
    out = _run(_chain(src))
    assert out["mode"] == "bilingual" and out["name"].endswith("_双语对照.docx")


def test_review_prefix_and_can_disable(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    src = _source_docx(tmp_path / "r.docx", tmp_path)
    payload = _chain(src)
    payload["statuses"] = ["missing_number"] + payload["statuses"][1:]

    out = _run({**payload, "mode": "bilingual"}, handle="rev")
    assert out["review"] == 1
    assert _texts(Document(out["path"]))[1].startswith("⚠️ ")

    out2 = _run({**payload, "mode": "bilingual", "mark_review": False}, handle="rev2")
    assert not _texts(Document(out2["path"]))[1].startswith("⚠️ ")


def test_mark_review_ignored_by_overlay(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    src = _source_docx(tmp_path / "o.docx", tmp_path)
    payload = _chain(src)
    payload["statuses"] = ["missing_number"] + payload["statuses"][1:]
    out = _run({**payload, "mode": "overlay"})
    assert out["review"] == 1
    assert not _texts(Document(out["path"]))[0].startswith("⚠️ ")  # 原位模式不加前缀


# ── 映射与容错 ───────────────────────────────────────────────────────────────
def test_index_map_dedup_path(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    src = tmp_path / "dup.docx"
    doc = Document()
    doc.add_paragraph("Repeated clause text.")
    doc.add_paragraph("Repeated clause text.")
    doc.save(src)

    payload = _chain(src)
    segments = [u["text"] for u in payload["units"]]
    assert len(segments) == 2
    unique = ["重复条款译文。"]
    payload["index_map"] = [0, 0]
    payload["translations"] = unique
    payload["statuses"] = ["ok"]
    out = _run({**payload, "mode": "overlay"})
    texts = _texts(Document(out["path"]))
    assert texts == ["重复条款译文。", "重复条款译文。"]


def test_missing_translation_keeps_original(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    src = _source_docx(tmp_path / "m.docx", tmp_path)
    payload = _chain(src)
    payload["translations"] = payload["translations"][:1]  # 只有第一条
    out = _run({**payload, "mode": "overlay"})
    assert out["skipped"] >= 1
    texts = _texts(Document(out["path"]))
    assert texts[0] == payload["translations"][0]
    assert texts[1].startswith("This agreement")  # 无译文 → 原样保留


def test_rejects_bad_inputs(tmp_path, monkeypatch):
    from core.errors import ToolDomainError

    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path / "data"))
    src = _source_docx(tmp_path / "e.docx", tmp_path)
    payload = _chain(src)

    with pytest.raises(ToolDomainError, match="未知渲染模式"):
        _run({**payload, "mode": "side-by-side"})
    with pytest.raises(ToolDomainError, match="units 为空"):
        _run({**payload, "units": []})
    with pytest.raises(ToolDomainError, match="translations 为空"):
        _run({**payload, "translations": []})
    with pytest.raises(ToolDomainError, match="文件不存在"):
        _run({**payload, "file": str(tmp_path / "nope.docx")})


def test_render_document_rejects_unknown_mode(tmp_path):
    from docx import Document

    from command_shared.docx_render import render_document

    with pytest.raises(ValueError, match="未知渲染模式"):
        render_document(Document(), units=[], ranges=[], index_map=None, date_maps=None,
                        translations=[], statuses=None, mode="nope")


def test_manifest_declares_expected_io():
    from core.protocol import ToolManifest

    manifest = ToolManifest.from_toml(
        Path(__file__).resolve().parents[1] / "tools" / "docx" / "render_translated" / "tool.toml")
    assert manifest.tool.id == "docx.render.translated"
    assert manifest.io.output_types == ["file.docx"]
    assert manifest.io.input_schema["properties"]["mode"]["default"] == "bilingual"
