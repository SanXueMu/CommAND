"""docx 翻译链离线端到端（无 LLM、无网络）：extract → classify → dedup → render。

只把「翻译」一步换成确定性的纯中文译文，用来锁定四件事：
1) 各步产物键与形状彼此咬合（units → ranges/col_classes → index_map/date_maps → translations）；
2) overlay 后**文档内不再有任何英文字母**（原文被真正替换，而非叠加）；
3) bilingual 保留原文并紧随插入中文段（表格格内对照）；
4) 源文件始终只读。
"""
from __future__ import annotations

import importlib
import re
import types
from pathlib import Path

import pytest

DRAW_XPATH = ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"


def _tool(name: str):
    return importlib.import_module(f"tools.{name}.main")


def _source(path: Path) -> Path:
    """标题 / 正文 / 空段 / 纯数字段 / 表格（文本列 + 数字列）。"""
    from docx import Document

    doc = Document()
    doc.add_heading("Audit Engagement Letter", level=1)
    doc.add_paragraph("This agreement is made between ZTH CPA and the Client.")
    doc.add_paragraph("")
    doc.add_paragraph("12345")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Audit fee"
    table.cell(1, 1).text = "8888"
    doc.save(str(path))
    return path


@pytest.fixture()
def chain(tmp_path):
    src = _source(tmp_path / "src.docx")
    ctx = types.SimpleNamespace(handle="chain")
    emit = lambda _event: None  # noqa: E731
    units = _tool("docx.extract_units").run({"file": str(src)}, ctx, emit)["units"]
    classified = _tool("table.classify_columns").run({"units": units}, ctx, emit)
    deduped = _tool("text.dedup_values").run({"segments": classified["segments"]}, ctx, emit)
    translations = [f"第{i + 1}段中文译文" for i in range(len(deduped["unique"]))]
    render = _tool("docx.render_translated")
    base = dict(
        file=str(src), units=units, ranges=classified["ranges"],
        col_classes=classified["col_classes"], index_map=deduped["index_map"],
        date_maps=deduped["date_maps"], translations=translations,
    )
    return src, base, render, ctx, emit


def test_chain_artifacts_fit_together(chain) -> None:
    _src, base, _render, _ctx, _emit = chain
    assert [(u["unit_id"], u["unit_type"]) for u in base["units"]] == [
        ("p0", "text"), ("p1", "text"), ("t4", "table")], "空段与纯数字段不成单元"
    assert sum(r["count"] for r in base["ranges"]) == 5, "2 个文字段 + 表格 3 格（数字列被排除）"
    text_counts = [r["count"] for u, r in zip(base["units"], base["ranges"]) if u["unit_type"] == "text"]
    assert text_counts == [1, 1], "text 单元恒为单段（回填按段写回）"
    assert base["col_classes"]["t4"][1]["class"] == "skip", "数字列应被排除"
    assert max(base["index_map"]) + 1 == len(base["translations"]), "index_map 指向去重后的唯一序列"


def _dump(path: str) -> tuple[list[str], list[str]]:
    from docx import Document

    doc = Document(path)
    return [p.text for p in doc.paragraphs], [c.text for t in doc.tables for r in t.rows for c in r.cells]


def test_overlay_replaces_source_entirely(chain) -> None:
    _src, base, render, ctx, emit = chain
    result = render.run({**base, "mode": "overlay"}, ctx, emit)
    paragraphs, cells = _dump(result["path"])
    assert result["inserted"] == 0 and result["replaced"] > 0 and result["skipped"] == 0
    assert not [t for t in paragraphs + cells if re.search(r"[A-Za-z]", t)], "overlay 后不应有任何英文残留"
    assert "12345" in paragraphs, "纯数字段保持原样"
    assert "8888" in cells, "数字列保持原样"
    from docx import Document

    assert Document(result["path"]).paragraphs[0].style.name.startswith("Heading"), "标题样式保留"


def test_bilingual_keeps_source_and_inserts_chinese(chain) -> None:
    _src, base, render, ctx, emit = chain
    result = render.run({**base, "mode": "bilingual"}, ctx, emit)
    paragraphs, cells = _dump(result["path"])
    source = "This agreement is made between ZTH CPA and the Client."
    index = paragraphs.index(source)
    assert paragraphs[index + 1].startswith("第") and "中文译文" in paragraphs[index + 1]
    assert "Audit fee" in "\n".join(cells) and "中文译文" in "\n".join(cells)
    assert result["replaced"] == 0 and result["inserted"] == sum(r["count"] for r in base["ranges"])


def test_source_file_never_modified(chain) -> None:
    src, base, render, ctx, emit = chain
    before = src.read_bytes()
    render.run({**base, "mode": "overlay"}, ctx, emit)
    render.run({**base, "mode": "bilingual"}, ctx, emit)
    assert src.read_bytes() == before, "源文件被改写"
    from docx import Document

    assert "This agreement" in Document(str(src)).paragraphs[1].text


def test_picture_run_survives_overlay(tmp_path) -> None:
    from docx import Document
    from docx.shared import Inches
    import pymupdf

    png = tmp_path / "px.png"
    pdf = pymupdf.open()
    page = pdf.new_page(width=20, height=20)
    page.get_pixmap(dpi=40).save(str(png))
    pdf.close()
    src = tmp_path / "pic.docx"
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Please refer to the appendix ")
    para.add_run().add_picture(str(png), width=Inches(0.3))
    doc.save(str(src))

    ctx = types.SimpleNamespace(handle="chain")
    emit = lambda _event: None  # noqa: E731
    units = _tool("docx.extract_units").run({"file": str(src)}, ctx, emit)["units"]
    classified = _tool("table.classify_columns").run({"units": units}, ctx, emit)
    deduped = _tool("text.dedup_values").run({"segments": classified["segments"]}, ctx, emit)
    result = _tool("docx.render_translated").run(
        {"file": str(src), "units": units, "ranges": classified["ranges"],
         "col_classes": classified["col_classes"], "index_map": deduped["index_map"],
         "date_maps": deduped["date_maps"], "translations": ["仅中文"], "mode": "overlay"},
        ctx, emit,
    )
    out = Document(result["path"])
    assert out.paragraphs[0].text == "仅中文"
    assert out.paragraphs[0]._p.findall(DRAW_XPATH), "图片 run 被清掉"
