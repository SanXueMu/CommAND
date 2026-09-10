"""B2 模板生成域 + 导出域工具测试。"""
import json
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook
from PIL import Image

import command_shared.chat as chat_mod
from core.errors import ToolDomainError
from tools.layout.analyze_pdf import main as analyze_pdf
from tools.records.export_csv import main as export_csv
from tools.records.export_xlsx import main as export_xlsx
from tools.spec.gen_textllm import main as gen_textllm
from tools.spec.gen_vl import main as gen_vl
from tools.spec.validate_ocrspec import main as validate_ocrspec

GOOD_SPEC = {
    "task_spec": {"fields": ["项目名称", "送审金额"], "record_mode": "page",
                  "rules": "金额保留两位", "template": "你是决算表识别专家",
                  "example": {"项目名称": "示例", "送审金额": "100.00"}},
    "postprocess": [{"name": "amount", "code": "def page_hook(fields, add_note):\n"
                    "    if not fields.get('送审金额'):\n"
                    "        add_note('送审金额缺失待审')"}],
    "view_spec": {"name": "决算", "columns": ["项目", "金额"],
                  "group": {"mode": "value", "field": "项目名称"},
                  "aggregates": [{"column": "项目", "op": "group_key", "field": "项目名称"},
                                 {"column": "金额", "op": "join_values", "field": "送审金额"}]},
}


class _Msg:
    def __init__(self, content):
        self.content = content


class _Choice:
    def __init__(self, content):
        self.message = _Msg(content)


class _Resp:
    def __init__(self, content):
        self.choices = [_Choice(content)]
        self.usage = None


def _run(module, payload, keys=None):
    ctx_keys = dict(keys or {})

    class Ctx:
        keys = ctx_keys

    events = []
    out = module.run(payload, Ctx(), events.append)
    return out, events


def _emit_ok(events):
    assert events and events[-1]["type"] == "progress"


# ---------- spec_gen 解析 ----------

def test_parse_spec_json_variants():
    from command_shared.spec_gen import parse_spec_json
    raw = '```json\n{"task_spec": {"fields": ["a"]}, "junk": 1}\n```'
    spec = parse_spec_json(raw)
    assert spec["task_spec"]["fields"] == ["a"]
    assert spec["postprocess"] == [] and spec["view_spec"] == {}
    with pytest.raises(ValueError):
        parse_spec_json("完全没有 JSON")
    with pytest.raises(ValueError):
        parse_spec_json('{"view_spec": {}}')


def test_build_spec_user_layout():
    from command_shared.spec_gen import build_spec_user
    text = build_spec_user("识别决算表", '{"page_count": 3}')
    assert "识别决算表" in text and "page_count" in text


# ---------- layout.analyze.pdf ----------

def _make_table_pdf(path: Path):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 90), "项目送审金额表", fontsize=16, fontname="china-s")
    x0, y0, x1, y1 = 60, 110, 460, 210
    col_x = [x0, (x0 + x1) / 2, x1]
    row_y = [y0, y0 + (y1 - y0) / 3, y0 + 2 * (y1 - y0) / 3, y1]
    for x in col_x:
        page.draw_line(fitz.Point(x, y0), fitz.Point(x, y1))
    for y in row_y:
        page.draw_line(fitz.Point(x0, y), fitz.Point(x1, y))
    cells = ["项目名称", "送审金额", "甲项目", "100.00", "乙项目", "50.00"]
    for index, text in enumerate(cells):
        row, col = divmod(index, 2)
        page.insert_text((col_x[col] + 4, row_y[row] + 22), text, fontsize=10, fontname="china-s")
    doc.save(path)
    doc.close()
    return path


def test_layout_analyze(tmp_path):
    pdf = _make_table_pdf(tmp_path / "t.pdf")
    out, events = _run(analyze_pdf, {"file": str(pdf)})
    layout = out["layout"]
    assert layout["page_count"] == 1 and layout["is_text_pdf"] is True
    assert json.loads(out["layout_json"])["page_count"] == 1
    assert layout["pages"][0]["tables"], "fitz find_tables 应检出表格"
    table = layout["pages"][0]["tables"][0]
    assert table["col_count"] >= 2 and table["row_count"] >= 2
    assert "项目送审金额表" in layout["pages"][0]["headers"]
    assert events == []  # 单页不足10页阈值，无进度事件


def test_layout_analyze_rejects_image(tmp_path):
    image = tmp_path / "a.png"
    Image.new("RGB", (40, 40)).save(image)
    with pytest.raises(ToolDomainError):
        _run(analyze_pdf, {"file": str(image)})


# ---------- spec.gen.textllm / gen_vl ----------

def test_gen_textllm(tmp_path, monkeypatch):
    import json as _json
    reply = _json.dumps(GOOD_SPEC, ensure_ascii=False)
    monkeypatch.setattr(chat_mod, "make_client", lambda key, base_url=None: object())
    monkeypatch.setattr(chat_mod, "call_chat",
                        lambda *a, **k: (reply, {"prompt_tokens": 1, "completion_tokens": 2}))
    out, _ = _run(gen_textllm, {"requirement": "识别送审金额", "layout_json": "{}"},
                  keys={"ds": {"api_key": "sk-x", "is_default": True}})
    assert out["task_spec"]["fields"] == ["项目名称", "送审金额"]
    assert out["usage"]["prompt_tokens"] == 1


def test_gen_textllm_requires_layout():
    with pytest.raises(ToolDomainError):
        _run(gen_textllm, {"requirement": "x"})


def test_gen_vl(tmp_path, monkeypatch):
    image = tmp_path / "s.png"
    Image.new("RGB", (60, 30)).save(image)
    import json as _json
    import tools.spec.gen_vl.main as vl_main
    monkeypatch.setattr(vl_main, "call_vl", lambda *a, **k: _json.dumps(GOOD_SPEC, ensure_ascii=False))
    out, _ = _run(gen_vl, {"file": str(image), "requirement": "识别签单要素"},
                  keys={"ds": {"api_key": "sk-x", "is_default": True}})
    assert out["sampled_pages"] == [1]
    assert out["view_spec"]["name"] == "决算"


# ---------- spec.validate.ocrspec ----------

def test_validate_good_and_bad():
    out, events = _run(validate_ocrspec, dict(GOOD_SPEC))
    assert out["valid"] is True and out["errors"] == []
    assert out["hooks_count"] == 1
    _emit_ok(events)

    bad = dict(GOOD_SPEC)
    bad["postprocess"] = [{"name": "h", "code": "def wrong_name():"}]
    bad["view_spec"] = {**GOOD_SPEC["view_spec"],
                        "aggregates": [{"column": "X", "op": "not_an_op", "field": "项目名称"}]}
    out, _ = _run(validate_ocrspec, bad)
    assert out["valid"] is False
    assert any("钩子不可编译" in e for e in out["errors"])
    assert any("未知聚合算子" in e for e in out["errors"])

    warn_case = dict(GOOD_SPEC)
    warn_case["view_spec"] = {**GOOD_SPEC["view_spec"], "columns": ["项目", "金额", "幽灵列"]}
    out, _ = _run(validate_ocrspec, warn_case)
    assert out["valid"] is True and any("幽灵列" in w for w in out["warnings"])


# ---------- exports ----------

ROWS = [{"项目": "甲", "金额": "100"}, {"项目": "乙", "金额": "50"}]


def test_export_csv(tmp_path, monkeypatch):
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path))
    out, events = _run(export_csv, {"columns": ["项目", "金额"], "rows": ROWS, "name": "决算/2026"})
    content = Path(out["file"]).read_bytes()
    assert content.startswith(b"\xef\xbb\xbf") and "项目,金额" in content.decode("utf-8-sig")
    assert out["rows_count"] == 2
    _emit_ok(events)


def test_export_xlsx_splits(tmp_path, monkeypatch):
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path))
    splits = [{"key": "送审明细", "columns": ["项目"], "rows": [{"项目": "丙"}]},
              {"key": "送审明细" * 8, "columns": ["项目"], "rows": []}]
    out, _ = _run(export_xlsx, {"columns": ["项目", "金额"], "rows": ROWS,
                                "splits": splits, "name": "多sheet", "sheet_name": "总表"})
    workbook = load_workbook(out["file"])
    assert workbook.sheetnames[0] == "总表"
    assert "送审明细" in workbook.sheetnames
    assert len(workbook.sheetnames) == 3 and len(workbook.sheetnames) == len(set(workbook.sheetnames))
    assert all(len(name) <= 31 for name in workbook.sheetnames)
    assert out["rows_count"] == 3


def test_export_requires_rows(tmp_path, monkeypatch):
    monkeypatch.setenv("COMMAND_DATA_DIR", str(tmp_path))
    with pytest.raises(ToolDomainError):
        _run(export_csv, {"columns": ["x"]})

