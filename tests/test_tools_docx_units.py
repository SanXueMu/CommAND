"""docx.extract.units / command_shared.docx_units 测试（含与 harvest 链的契约对齐）。"""
from __future__ import annotations

import importlib
import types
from pathlib import Path

import pytest

from command_shared.docx_units import (
    body_index_map, extract_units, heading_level, is_toc_paragraph, iter_body,
)


def _tool():
    return importlib.import_module("tools.docx.extract_units.main")


def _field_paragraph(document, label: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = document.add_paragraph()
    paragraph.add_run(label)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    return paragraph


def _sample_docx(path: Path) -> Path:
    """标题 + 正文 + 空段 + 纯数字段 + 域段 + 表格 + 尾段（body 序号 0..6）。"""
    from docx import Document

    doc = Document()
    doc.add_heading("Audit Engagement Letter", level=1)
    doc.add_paragraph("This agreement is made between ZTH CPA and the Client.")
    doc.add_paragraph("")
    doc.add_paragraph("2026-01-01")
    _field_paragraph(doc, "Page ")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Audit fee"
    table.cell(1, 1).text = "123456"
    doc.add_paragraph("Signature: Zhang Yi")
    doc.save(path)
    return path


def _load(path: Path):
    from docx import Document

    return Document(str(path))


# ── 提取 ─────────────────────────────────────────────────────────────────────
def test_units_order_ids_and_skips(tmp_path):
    src = _sample_docx(tmp_path / "s.docx")
    result = extract_units(_load(src))
    units, stats = result["units"], result["stats"]

    # 正文流：p0 标题 / p1 正文 / p2 空段(跳) / p3 纯数字(跳) / p4 域段(跳) / t5 表格 / p6 尾段
    assert [u["unit_id"] for u in units] == ["p0", "p1", "t5", "p6"]
    assert [u["unit_type"] for u in units] == ["text", "text", "table", "text"]
    assert stats["paragraphs"] == 6 and stats["tables"] == 1 and stats["units"] == 4
    assert stats["skipped_empty"] == 1 and stats["skipped_number"] == 1
    assert stats["skipped_field"] == 1 and stats["text_chars"] > 0

    title = units[0]
    assert title["meta"]["body_index"] == 0 and title["meta"]["heading"] == 1
    assert title["meta"]["style"].startswith("Heading")
    assert "ZTH CPA" in units[1]["text"]


def test_table_unit_rows_and_meta(tmp_path):
    units = extract_units(_load(_sample_docx(tmp_path / "t.docx")))["units"]
    table = next(u for u in units if u["unit_type"] == "table")
    assert table["unit_id"] == "t5"
    assert table["rows"] == [["Item", "Amount"], ["Audit fee", "123456"]]
    assert table["meta"]["rows"] == 2 and table["meta"]["cols"] == 2


def test_body_index_map_aligns_with_iter_body(tmp_path):
    doc = _load(_sample_docx(tmp_path / "m.docx"))
    mapping = body_index_map(doc)
    assert set(mapping) == {"p0", "p1", "p2", "p3", "p4", "t5", "p6"}
    assert [kind for _, kind, _ in iter_body(doc)] == ["p", "p", "p", "p", "p", "tbl", "p"]
    assert mapping["p0"].text.startswith("Audit Engagement")
    assert len(mapping["t5"].rows) == 2


def test_switches_include_tables_and_skip_numbers(tmp_path):
    doc = _load(_sample_docx(tmp_path / "sw.docx"))
    no_tables = extract_units(doc, include_tables=False)["units"]
    assert all(u["unit_type"] == "text" for u in no_tables)
    with_numbers = extract_units(doc, skip_numbers=False)["units"]
    assert any(u.get("text") == "2026-01-01" for u in with_numbers)


def test_toc_style_detected():
    class _Style:
        name = "TOC 1"

    fake = types.SimpleNamespace(style=_Style())
    assert is_toc_paragraph(fake) is True
    assert heading_level(types.SimpleNamespace(style=types.SimpleNamespace(name="Heading 3"))) == 3
    assert heading_level(types.SimpleNamespace(style=types.SimpleNamespace(name="Normal"))) is None


# ── 与 harvest 链的契约对齐（关键：既有 docx 回填器要求 text 单元 count == 1）────────
def test_units_harvest_contract(tmp_path):
    from command_shared.table import classify_columns, harvest_units, unit_cells

    units = extract_units(_load(_sample_docx(tmp_path / "h.docx")))["units"]
    import re

    col_classes = {u["unit_id"]: classify_columns(u["rows"], re.compile(r"^[\d\s/\-:.,()%¥$€£#&+＋]+$"))
                   for u in units if u["unit_type"] == "table"}
    harvested = harvest_units(units, col_classes)

    assert harvested["segments"][0] == "Audit Engagement Letter"
    assert any("Audit fee" in s for s in harvested["segments"])
    # 每 text 单元恰好一段（既有 docx.render.bilingual 的 index_map[range.start] 语义）
    for unit, rng in zip(units, harvested["ranges"]):
        assert rng["unit_id"] == unit["unit_id"]
        assert rng["count"] >= 1
        if unit["unit_type"] == "text":
            assert rng["count"] == 1
    # 表格段数 == 逐格收割数（表头字母格 + 非 skip 列数据格）
    table = next(u for u in units if u["unit_type"] == "table")
    rng = next(r for r in harvested["ranges"] if r["unit_id"] == table["unit_id"])
    assert rng["count"] == len(unit_cells(table["rows"], col_classes[table["unit_id"]]))


# ── 工具接线 ─────────────────────────────────────────────────────────────────
def test_tool_run_wiring(tmp_path):
    src = _sample_docx(tmp_path / "run.docx")
    events = []
    out = _tool().run({"file": str(src)}, types.SimpleNamespace(handle="h"), events.append)
    assert out["kind"] == "docx_units" and out["units"] and len(out["file_hash"]) == 64
    assert out["stats"]["units"] == len(out["units"])
    assert events and events[-1]["phase"] == "extracted"


def test_tool_pauses_on_legacy_doc_and_rejects_missing(tmp_path):
    """旧版 .doc（OLE2 魔数）→ ToolPauseError（任务暂停等人工另存为 .docx），不是普通失败。"""
    from core.errors import ToolPauseError, ToolDomainError

    old = tmp_path / "old.doc"
    old.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(ToolPauseError, match="旧版 .doc"):
        _tool().run({"file": str(old)}, types.SimpleNamespace(handle="h"), lambda e: None)

    # 魔数优先于后缀：误命名为 .docx 的二进制 .doc 同样暂停
    fake = tmp_path / "actually_old.docx"
    fake.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(ToolPauseError, match="旧版 .doc"):
        _tool().run({"file": str(fake)}, types.SimpleNamespace(handle="h"), lambda e: None)

    with pytest.raises(ToolDomainError, match="文件不存在"):
        _tool().run({"file": str(tmp_path / "nope.docx")}, types.SimpleNamespace(handle="h"), lambda e: None)

    from docx import Document

    empty = tmp_path / "empty.docx"
    doc = Document()
    doc.add_paragraph("")
    doc.add_paragraph("1234")
    doc.save(empty)
    with pytest.raises(ToolDomainError, match="未提取到可翻译内容"):
        _tool().run({"file": str(empty)}, types.SimpleNamespace(handle="h"), lambda e: None)


def test_manifest_declares_expected_io():
    from core.protocol import ToolManifest

    manifest = ToolManifest.from_toml(
        Path(__file__).resolve().parents[1] / "tools" / "docx" / "extract_units" / "tool.toml")
    assert manifest.tool.id == "docx.extract.units"
    assert manifest.runtime.entry == "main.py:run"
    assert manifest.io.input_types == ["file.docx"]
    assert manifest.io.output_types == ["docx.units"]
