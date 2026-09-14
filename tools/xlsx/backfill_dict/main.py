"""译文回插原工作簿 → 中文版 xlsx（translee backfill xlsx 半边移植，纯数据 join）。

两种版式（`layout`）：
- `sheets`（默认，交付口径）：在**原文件副本**上，每个原 sheet 之后插入 `<原sheet名>_翻译结果`，
  译文按原单元格坐标回填（原 sheet 一字不动），产出「原文｜译文」并排的工作簿；**不含对照字典**。
- `dict`（旧行为，跨文件汇总译文用）：另建工作簿，每个表单单元一张 sheet + 末尾「对照字典」。

对齐链：单元 → ranges.start → unit_cells 收割序 → index_map → unique 下标 → 译文。
对应关系全部来自显式输入数据（零隐式共享知识）；review 格保留原文加 ⚠️ 前缀。
源文件始终只读：绝不写回原文件。
"""
from __future__ import annotations

import os
import re
from pathlib import Path

from command_shared.table import unit_cells

_REVIEW_PREFIX = "⚠️ "
_SHEET_UNSAFE_RE = re.compile(r"[\[\]:*?/\\]")
_RESULT_SUFFIX = "_翻译结果"
_MAX_TITLE = 31
LAYOUT_SHEETS = "sheets"
LAYOUT_DICT = "dict"


def _safe_sheet_title(title: str) -> str:
    return _SHEET_UNSAFE_RE.sub("_", title)[:31] or "Sheet"


def _result_title(title: str, taken: set[str]) -> str:
    """`<原sheet名>_翻译结果`，兼顾 31 字符上限与重名。"""
    base = _SHEET_UNSAFE_RE.sub("_", title)
    root = base[: _MAX_TITLE - len(_RESULT_SUFFIX)] or "Sheet"
    cand = root + _RESULT_SUFFIX
    n = 2
    while cand in taken:
        tail = f"_{n}"
        cand = root[: _MAX_TITLE - len(_RESULT_SUFFIX) - len(tail)] + _RESULT_SUFFIX + tail
        n += 1
    return cand


def _cell_result_map(unit: dict, col_classes: list[dict], start: int,
                     index_map: list[int], date_maps: list[dict],
                     translations: list[str], statuses: list[str]) -> dict[tuple[int, int], dict]:
    """单元 → {(row, col): {final, status}}（收割序与 classify 逐位一致）。"""
    result: dict[tuple[int, int], dict] = {}
    rows = unit["rows"]
    for j, (row_idx, col_idx) in enumerate(unit_cells(rows, col_classes)):
        g = start + j
        u = index_map[g]
        entry = {"status": statuses[u], "dates": date_maps[g]}
        if statuses[u] == "ok":
            final = translations[u]
            for placeholder, value in entry["dates"].items():
                final = final.replace(placeholder, value)
            entry["final"] = final
        result[(row_idx, col_idx)] = entry
    return result


def _render_row(row: list[str], cell_map: dict[tuple[int, int], dict], row_idx: int) -> list[str]:
    out = []
    for col, cell in enumerate(row):
        entry = cell_map.get((row_idx, col))
        if entry is None:
            out.append(cell)
        elif entry["status"] == "ok":
            out.append(entry["final"])
        else:
            out.append(_REVIEW_PREFIX + cell)
    return out


def _apply_cells(ws, rows: list[list[str]], cell_map: dict[tuple[int, int], dict]) -> int:
    """回填到**已存在的单元格**（行列 +1 对齐 A1 起点），保留原样式；返回改写格数。"""
    changed = 0
    for (row_idx, col_idx), entry in cell_map.items():
        if entry["status"] == "ok":
            value = entry["final"]
        else:
            original = rows[row_idx][col_idx] if col_idx < len(rows[row_idx]) else ""
            value = _REVIEW_PREFIX + original
        ws.cell(row=row_idx + 1, column=col_idx + 1).value = value
        changed += 1
    return changed


def _render_sheets(src: Path, units: list[dict], range_by_unit: dict, col_classes: dict,
                   index_map: list[int], date_maps: list[dict],
                   translations: list[str], statuses: list[str], out_path: Path) -> dict:
    """在原文件副本上插入译文 sheet（每个原 sheet 之后一张），原 sheet 不动。"""
    from openpyxl import load_workbook

    wb = load_workbook(src)  # 副本在内存中改写，源文件只读
    unit_by_title = {u["unit_id"]: u for u in units if u.get("unit_type") == "table"}
    changed = 0
    made = 0
    for ws in list(wb.worksheets):
        unit = unit_by_title.get(ws.title)
        if unit is None:
            continue
        rng = range_by_unit.get(ws.title)
        if rng is None:
            continue
        cell_map = _cell_result_map(unit, col_classes.get(ws.title, []), rng["start"],
                                    index_map, date_maps, translations, statuses)
        copy = wb.copy_worksheet(ws)  # 样式随副本保留
        copy.title = _result_title(ws.title, set(wb.sheetnames))
        changed += _apply_cells(copy, unit["rows"], cell_map)
        made += 1
        # 放到原 sheet 正后方（copy_worksheet 默认追加在末尾）
        target = wb.sheetnames.index(ws.title) + 1
        current = wb.sheetnames.index(copy.title)
        wb.move_sheet(copy, offset=target - current)

    wb.save(out_path)
    return {"sheets_made": made, "cells_changed": changed}


def _render_dict(units: list[dict], range_by_unit: dict, col_classes: dict, segments: list[str],
                 index_map: list[int], date_maps: list[dict], translations: list[str],
                 statuses: list[str], out_path: Path) -> dict:
    """旧行为：新建工作簿，每单元一张 sheet + 末尾「对照字典」。"""
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    dict_rows: list[tuple[str, str, str, int]] = []
    seen_titles: set[str] = set()
    for unit in units:
        if unit.get("unit_type") != "table":
            continue
        unit_id = unit["unit_id"]
        rng = range_by_unit.get(unit_id)
        if rng is None:
            continue
        cell_map = _cell_result_map(unit, col_classes.get(unit_id, []), rng["start"],
                                    index_map, date_maps, translations, statuses)
        rows = unit["rows"]

        title = _safe_sheet_title(unit_id)
        n = 2
        while title in seen_titles:
            title = f"{_safe_sheet_title(unit_id)[:28]}_{n}"
            n += 1
        seen_titles.add(title)
        ws = wb.create_sheet(title)
        for row_idx, row in enumerate(rows):
            ws.append(_render_row(row, cell_map, row_idx))

        counts: dict[int, int] = {}
        for g in index_map:
            counts[g] = counts.get(g, 0) + 1
        for j in range(rng["count"]):
            g = rng["start"] + j
            u = index_map[g]
            final = translations[u]
            for placeholder, value in date_maps[g].items():
                final = final.replace(placeholder, value)
            dict_rows.append((segments[g], final, statuses[u], counts[u]))

    ws_dict = wb.create_sheet("对照字典")
    ws_dict.append(["原文（日期归一前采样）", "译文", "状态", "引用次数"])
    for sample, translated, status, count in sorted(dict_rows, key=lambda r: -r[3]):
        ws_dict.append([sample, translated, status, count])

    wb.save(out_path)
    return {"dict_rows": len(dict_rows)}


def run(input: dict, ctx, emit) -> dict:
    units = input["units"]
    col_classes = input["col_classes"]
    range_by_unit = {r["unit_id"]: r for r in input["ranges"]}
    segments = input["segments"]
    index_map = input["index_map"]
    date_maps = input["date_maps"]
    translations = input["translations"]
    statuses = input["statuses"]

    data_dir = Path(os.environ.get("COMMAND_DATA_DIR", "data"))
    out_dir = data_dir / "outputs" / ctx.handle
    out_dir.mkdir(parents=True, exist_ok=True)

    src = Path(input["file"]) if input.get("file") else None
    stem = src.stem if src else "output"
    out_name = input.get("out_name") or f"{stem}_中文.xlsx"
    out_path = out_dir / out_name

    layout = input.get("layout") or LAYOUT_SHEETS
    if layout == LAYOUT_DICT or not (src and src.is_file()):
        # dict 模式，或源文件不可得（历史数据/纯夹具）时回落旧行为
        stats = _render_dict(units, range_by_unit, col_classes, segments, index_map,
                             date_maps, translations, statuses, out_path)
        layout = LAYOUT_DICT
    else:
        stats = _render_sheets(src, units, range_by_unit, col_classes, index_map,
                               date_maps, translations, statuses, out_path)

    emit({"phase": "backfilled", "path": str(out_path), "layout": layout, **stats})
    return {"path": str(out_path), "name": out_name, "layout": layout, **stats}
