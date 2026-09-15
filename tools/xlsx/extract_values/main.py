"""xlsx → 每 sheet 一个 table 单元（translee extract_xlsx 移植，零翻译语义）。

大表内存策略：read_only 流式逐 sheet 提取，单元格全部转字符串；去尾部全空行。
"""
from __future__ import annotations

import hashlib
from pathlib import Path


def _file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(1048576):
            h.update(chunk)
    return h.hexdigest()


def _format_cell(value) -> str:
    """单元格 → 字符串：空值→""，整值浮点去 .0。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def run(input: dict, ctx, emit) -> dict:
    path = Path(input["file"])
    if not path.is_file():
        from core.errors import ToolDomainError

        raise ToolDomainError(f"文件不存在: {path}")

    from openpyxl import load_workbook

    from command_shared.xls_convert import is_xls, xls_to_xlsx

    fhash = _file_hash(path)
    # 老版 .xls（OLE2）：openpyxl 不支持，先转 .xlsx 再走原链路（值为准）
    conv_tmp = xls_to_xlsx(path) if is_xls(path) else None
    read_path = conv_tmp or path
    units = []
    wb = load_workbook(read_path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = [
                [_format_cell(cell) for cell in row]
                for row in ws.iter_rows(values_only=True)
            ]
            while rows and not any(rows[-1]):
                rows.pop()
            if not rows:
                continue
            units.append({"unit_id": ws.title, "unit_type": "table", "rows": rows})
    finally:
        wb.close()
        if conv_tmp is not None:
            conv_tmp.unlink(missing_ok=True)

    emit({"phase": "extracted", "units": len(units)})
    return {"file": str(path), "file_hash": fhash, "kind": "xlsx", "units": units}
