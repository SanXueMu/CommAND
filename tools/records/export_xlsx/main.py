"""视图行导出 xlsx：主表 + 可选 splits 多 sheet（每 sheet 列集独立、空列已裁）。

openpyxl（已有依赖）；sheet 名 ≤31 字符且去非法字符，重名自动加序号。
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from core.errors import ToolDomainError

MAX_SHEET_NAME = 31

# AO4：纯数字串 → 数值单元格（Excel 才能按数值排序/求和；文本「10」会排在「9」前）
_NUM_RE = re.compile(r"^-?\d+(\.\d+)?$")
_MAX_NUM_DIGITS = 12  # 超长纯数字（流水号/证件号）保留文本，避免精度丢失与科学计数


def _cell(value):
    """纯数字串写数值；其余（含前导零串、超长数字、中文/日期）一律保留原文。"""
    if not isinstance(value, str):
        return value
    text = value.strip()
    digits = text.replace("-", "").replace(".", "")
    if not text or len(digits) > _MAX_NUM_DIGITS:
        return value
    if len(text) > 1 and text[0] == "0" and "." not in text:
        return value  # 前导零（如凭证号 007）保留文本，不丢位
    if _NUM_RE.match(text):
        return float(text) if "." in text else int(text)
    return value


def _sheet_title(raw: str, used: set[str]) -> str:
    base = re.sub(r'[\\/*?:\[\]]+', "_", (raw or "Sheet").strip())[:MAX_SHEET_NAME] or "Sheet"
    title, index = base, 2
    while title in used:
        suffix = f"_{index}"
        title = base[:MAX_SHEET_NAME - len(suffix)] + suffix
        index += 1
    used.add(title)
    return title


def _write_sheet(worksheet, columns: list[str], rows: list) -> None:
    worksheet.append(columns)
    for row in rows:
        worksheet.append([_cell(row.get(c, "")) if isinstance(row, dict) else row for c in columns])
    for index, column in enumerate(columns, start=1):
        width = max(10, min(48, max((len(str(row.get(column, ""))) for row in rows), default=10) + 4))
        worksheet.column_dimensions[get_column_letter(index)].width = width


def run(input: dict, ctx, emit) -> dict:
    rows = input.get("rows")
    columns = input.get("columns") or []
    if rows is None:
        raise ToolDomainError("rows 必填（records.view.query 产出）")

    data_dir = Path(__import__("os").environ.get("COMMAND_DATA_DIR", "data"))
    out_dir = data_dir / "outputs" / "exports"
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{re.sub(r'[\\/:*?\"<>|]+', '_', (input.get('name') or '视图导出').strip() or '视图导出')}.xlsx"

    workbook = Workbook()
    used: set[str] = set()
    _write_sheet(workbook.active, columns, rows)
    workbook.active.title = _sheet_title(input.get("sheet_name") or "视图", used)

    sheets = [workbook.active.title]
    total_rows = len(rows)
    for split in input.get("splits") or []:
        worksheet = workbook.create_sheet(_sheet_title(str(split.get("key", "表")), used))
        split_columns = split.get("columns") or []
        split_rows = split.get("rows") or []
        _write_sheet(worksheet, split_columns, split_rows)
        sheets.append(worksheet.title)
        total_rows += len(split_rows)

    workbook.save(target)
    emit({"type": "progress", "phase": "export",
          "message": f"导出 {len(sheets)} sheet / {total_rows} 行 → {target.name}"})
    return {"file": str(target), "sheets": sheets, "rows_count": total_rows}
